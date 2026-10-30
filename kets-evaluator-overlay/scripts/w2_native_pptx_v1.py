"""Editable PowerPoint twin of the W2 weekly letter: paragraph text boxes and real tables.

Design Ref: user 2026-10-12 ("ppt에 표로 되어 있는 것도 다 개체로 되어 있어서 조절하기 어려운데, 그냥 표로
만들거나 개체를 좀 더 수정하기 편하게 해줘" / "단어 단위로 되어 있으니 텍스트 수정하기가 너무 어려워").
The previous editable deck put every word in its own text box over a picture. This module records
the drawing calls of the real PDF build (labels, paragraphs, tables, article cards, stat boxes,
charts, rectangles, lines, images, links) without changing what is drawn, then rebuilds each page:
one text box per paragraph or label, a PowerPoint table per data table, one shape per article card
(title, bullets and source link in its text frame), one shape per stat box, and the chart as a
picture cropped from the PDF render. Every text, number and link comes from the recorded calls of
the verified PDF; nothing is retyped.
Font (user 2026-10-12 "글꼴도 깨지지 않도록하는 기본 설정 값으로 바꿔줘"): the deck uses Malgun Gothic,
the default Korean font of Windows and Office, so it opens the same on any office PC. Malgun is about
12% wider than the PDF's IBM Plex Sans KR for Hangul, so each text item is set at the size whose
Malgun width matches the PDF width (line breaks then match); Plex Medium maps to bold from 10pt up.
Plan SC: PowerPoint rounds exact line spacing to whole points; with Malgun Gothic the first baseline
sits at 0.755 of the line, or 0.82 when the line is under 1.205 x the font size (measured 2026-10-12);
boxes are placed so baselines match the PDF. Runs carry lang="ko-KR": with en-US PowerPoint breaks
Hangul inside words (eaLnBrk alone does not help); with ko-KR it breaks at spaces like the PDF.
"""

from __future__ import annotations

import html
import inspect
import math
import re
import sys
from collections.abc import Callable, Sequence
from pathlib import Path
from typing import Any

RUNTIME = (Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python"
           / "Lib/site-packages")
if RUNTIME.exists() and str(RUNTIME) not in sys.path:
    sys.path.append(str(RUNTIME))

# isort: off
from lxml import etree  # type: ignore[import-untyped]  # noqa: E402
from pptx import Presentation  # type: ignore[import-not-found]  # noqa: E402
from pptx.dml.color import RGBColor  # type: ignore[import-not-found]  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # type: ignore[import-not-found]  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # type: ignore[import-not-found]  # noqa: E402
from pptx.oxml.ns import qn  # type: ignore[import-not-found]  # noqa: E402
from pptx.util import Pt  # type: ignore[import-not-found]  # noqa: E402
# isort: on

FONT = "Malgun Gothic"  # 맑은 고딕: installed with Windows and Office
MALGUN_FILES = {False: Path("C:/Windows/Fonts/malgun.ttf"), True: Path("C:/Windows/Fonts/malgunbd.ttf")}
FIT_MARGIN = 0.985  # Hangul/Latin mix differs line by line; keep each line a little inside the PDF width
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
SLACK = 2.0  # same font advances as the PDF; a little room for rounding only


def _hex(color: Any) -> str:
    if color is None:
        return "000000"
    if isinstance(color, str):
        return color.lstrip("#").upper()
    parts = color[:3] if isinstance(color, (tuple, list)) else (color.red, color.green, color.blue)
    return "".join(f"{round(float(c) * 255):02X}" for c in parts)


def _lead(value: float) -> float:
    return float(max(1, round(value)))


# ---------------------------------------------------------------- recording
def _maker(shape: Callable[[str, Any, Callable[..., dict[str, Any]]], Any], kind: str,
           unpack: Callable[..., dict[str, Any]]) -> Callable[[Any], Any]:
    def make(original: Any) -> Any:
        return shape(kind, original, unpack)
    return make


class Recorder:
    """Records the drawing calls of one PDF build, page by page."""

    def __init__(self, page_height: float) -> None:
        self.H = page_height
        self.pages: list[list[dict[str, Any]]] = [[]]
        self.stack: list[dict[str, Any]] = []
        self._undo: list[tuple[Any, str, Any]] = []
        self._in_label = 0

    def add(self, item: dict[str, Any]) -> None:
        (self.stack[-1]["items"] if self.stack else self.pages[-1]).append(item)

    def _patch(self, owner: Any, name: str, make: Callable[[Any], Any]) -> None:
        original = getattr(owner, name)
        setattr(owner, name, make(original))
        self._undo.append((owner, name, original))

    def uninstall(self) -> None:
        for owner, name, original in reversed(self._undo):
            setattr(owner, name, original)
        self._undo.clear()
        if self.pages and not self.pages[-1]:
            self.pages.pop()

    def group(self, owner: Any, name: str, kind: str) -> None:
        """Record calls of a drawing helper as one group holding its inner drawing calls."""
        if not hasattr(owner, name):
            return
        signature = inspect.signature(getattr(owner, name))

        def make(original: Any) -> Any:
            def inner(*args: Any, **kwargs: Any) -> Any:
                bound = signature.bind(*args, **kwargs)
                bound.apply_defaults()
                group = {"kind": kind, "args": {k: v for k, v in bound.arguments.items() if k != "pdf"},
                         "items": []}
                self.add(group)
                self.stack.append(group)
                try:
                    result = original(*args, **kwargs)
                finally:
                    self.stack.pop()
                group["result"] = result
                return result
            return inner
        self._patch(owner, name, make)

    def install(self, base: Any, groups: Sequence[tuple[Any, str, str]]) -> None:
        from reportlab import rl_config  # type: ignore[import-untyped]
        from reportlab.pdfgen.canvas import Canvas  # type: ignore[import-untyped]
        from reportlab.platypus.paragraph import Paragraph  # type: ignore[import-untyped]
        rec, H = self, self.H

        def label(original: Any) -> Any:
            def inner(pdf: Any, value: str, x: float, top: float, size: float, **kw: Any) -> Any:
                face = kw.get("face", "Regular")
                rec.add({"kind": "label", "text": value, "x": x, "top": top, "size": size, "face": face,
                         "color": _hex(kw.get("color", base.INK)), "baseline": top + size,
                         "width": base.text_width(value, size, face), "lines": 1,
                         "leading": size * 1.2})
                rec._in_label += 1
                try:
                    return original(pdf, value, x, top, size, **kw)
                finally:
                    rec._in_label -= 1
            return inner
        self._patch(base, "label", label)

        def draw_on(original: Any) -> Any:
            def inner(para: Any, canvas: Any, x: float, y: float, _sW: float = 0) -> Any:
                st = para.style
                offset = st.fontSize if getattr(rl_config, "paraFontSizeHeightOffset", 1) else getattr(
                    para.blPara, "ascent", st.fontSize)
                top = H - (y + para.height)
                rec.add({"kind": "para", "markup": para.text, "x": x, "top": top, "width": para.width,
                         "height": para.height, "size": st.fontSize, "leading": st.leading,
                         "face": st.fontName, "color": _hex(st.textColor), "left_indent": st.leftIndent,
                         "first_indent": st.firstLineIndent, "lines": len(para.blPara.lines),
                         "baseline": top + offset})
                return original(para, canvas, x, y, _sW)
            return inner
        self._patch(Paragraph, "drawOn", draw_on)

        def shape(kind: str, original: Any, unpack: Callable[..., dict[str, Any]]) -> Any:
            def inner(canvas: Any, *args: Any, **kwargs: Any) -> Any:
                if not rec._in_label:
                    item = unpack(*args, **kwargs)
                    item.update(kind=kind, fill_color=_hex(canvas._fillColorObj),
                                stroke_color=_hex(canvas._strokeColorObj), line_width=canvas._lineWidth)
                    rec.add(item)
                return original(canvas, *args, **kwargs)
            return inner

        def rect(x: float, y: float, width: float, height: float, stroke: int = 1, fill: int = 0,
                 **_: Any) -> dict[str, Any]:
            return {"x": x, "top": H - y - height, "w": width, "h": height, "fill": fill, "stroke": stroke}

        def round_rect(x: float, y: float, width: float, height: float, radius: float, stroke: int = 1,
                       fill: int = 0, **_: Any) -> dict[str, Any]:
            return {"x": x, "top": H - y - height, "w": width, "h": height, "radius": radius,
                    "fill": fill, "stroke": stroke}

        def line(x1: float, y1: float, x2: float, y2: float, **_: Any) -> dict[str, Any]:
            return {"x1": x1, "y1": H - y1, "x2": x2, "y2": H - y2}

        def circle(x: float, y: float, r: float, stroke: int = 1, fill: int = 0, **_: Any) -> dict[str, Any]:
            return {"x": x - r, "top": H - y - r, "w": 2 * r, "h": 2 * r, "fill": fill, "stroke": stroke}

        def image(image: Any, x: float, y: float, width: float | None = None, height: float | None = None,
                  *_: Any, **__: Any) -> dict[str, Any]:
            return {"path": str(image), "x": x, "top": H - y - (height or 0), "w": width, "h": height}

        def link_url(url: str, rect: Sequence[float], *_: Any, **__: Any) -> dict[str, Any]:
            x1, y1, x2, y2 = rect
            return {"url": url, "box": (min(x1, x2), H - max(y1, y2), max(x1, x2), H - min(y1, y2))}

        def link_absolute(contents: Any, destinationname: str, Rect: Sequence[float] | None = None,
                          *_: Any, **__: Any) -> dict[str, Any]:
            x1, y1, x2, y2 = Rect or (0, 0, 0, 0)
            return {"dest": destinationname, "box": (min(x1, x2), H - max(y1, y2), max(x1, x2), H - min(y1, y2))}

        for name, kind, unpack in (("rect", "rect", rect), ("roundRect", "roundrect", round_rect),
                                   ("line", "line", line), ("circle", "circle", circle),
                                   ("drawImage", "image", image), ("linkURL", "link", link_url),
                                   ("linkAbsolute", "jump", link_absolute)):
            self._patch(Canvas, name, _maker(shape, kind, unpack))

        def show_page(original: Any) -> Any:
            def inner(canvas: Any) -> Any:
                result = original(canvas)
                rec.pages.append([])
                return result
            return inner
        self._patch(Canvas, "showPage", show_page)
        for owner, name, kind in groups:
            self.group(owner, name, kind)


# ---------------------------------------------------------------- deck writing
def _runs(markup: str) -> list[tuple[str, bool]]:
    out, bold = [], False
    for piece in re.split(r"(</?b>)", markup):
        if piece == "<b>":
            bold = True
        elif piece == "</b>":
            bold = False
        elif piece:
            out.append((html.unescape(re.sub(r"<[^>]+>", "", piece)), bold))
    return out


def _bold_face(face: str, size: float) -> bool:
    return face == "Bold" or (face == "Medium" and size >= 10)


def _malgun(bold: bool) -> str | None:
    from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
    from reportlab.pdfbase.ttfonts import TTFont  # type: ignore[import-untyped]

    name = "W2NativeMalgunBold" if bold else "W2NativeMalgun"
    if name not in pdfmetrics.getRegisteredFontNames():
        if not MALGUN_FILES[bold].is_file():
            return None
        pdfmetrics.registerFont(TTFont(name, str(MALGUN_FILES[bold])))
    return name


BULLET_PREFIX = "-&nbsp;"  # the PDF draws summary bullets as '- ' with a hanging indent


def _is_bullet(item: dict[str, Any]) -> bool:
    return item["kind"] == "para" and str(item["markup"]).startswith(BULLET_PREFIX)


def _pieces(item: dict[str, Any]) -> list[tuple[str, bool]]:
    """Text runs; a summary bullet loses its typed '- ' (the deck uses a real bullet)."""
    if item["kind"] != "para":
        return [(item["text"], False)]
    markup = str(item["markup"])
    return _runs(markup[len(BULLET_PREFIX):] if _is_bullet(item) else markup)


def _deck_size(item: dict[str, Any]) -> float:
    """Malgun size whose width equals the PDF width of the same text (never larger than the PDF size)."""
    from reportlab.pdfbase import pdfmetrics

    size = float(item["size"])
    pdf_width = deck_width = 0.0
    try:
        for text, bold in _pieces(item):
            name = _malgun(bold or _bold_face(item["face"], size))
            if name is None:
                return size
            pdf_width += pdfmetrics.stringWidth(text, "Bold" if bold else item["face"], size)
            deck_width += pdfmetrics.stringWidth(text, name, size)
    except Exception:  # face not registered (tests, other builders): keep the PDF size
        return size
    if deck_width <= 0:
        return size
    return math.floor(size * min(1.0, pdf_width / deck_width * FIT_MARGIN) * 10) / 10


def _metrics(item: dict[str, Any]) -> tuple[float, float, float]:
    """(deck font size, exact line spacing, first-baseline offset inside the line)."""
    if "_metrics" not in item:
        size = _deck_size(item)
        lead = _lead(item["leading"]) if item["kind"] == "para" else _lead(1.3 * size)
        offset = (0.82 if lead < 1.205 * size else 0.755) * lead
        item["_metrics"] = (size, lead, offset)
    metrics: tuple[float, float, float] = item["_metrics"]
    return metrics


def _font(run: Any, size: float, color: str, bold: bool = False, underline: bool = False) -> None:
    name = FONT
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ko-KR")
    rpr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rpr, qn(tag))
        node.set("typeface", name)


def _frame(frame: Any, wrap: bool) -> None:
    frame.word_wrap = wrap
    frame.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, 0)
    frame.vertical_anchor = MSO_ANCHOR.TOP


def _ppr(paragraph: Any) -> Any:
    paragraph.alignment = PP_ALIGN.LEFT  # shapes default to centred text
    node = paragraph._p.get_or_add_pPr()
    node.set("eaLnBrk", "0")  # break Korean lines at spaces like the PDF
    return node


def _fill_paragraphs(frame: Any, items: Sequence[dict[str, Any]], left: float, first_top: float) -> float:
    """Write recorded text items as paragraphs; returns the bottom of the last line."""
    bottom = first_top
    for i, item in enumerate(items):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        size, lead, offset = _metrics(item)
        paragraph.line_spacing = Pt(lead)
        if i:
            gap = (item["baseline"] - offset) - bottom
            paragraph.space_before = Pt(max(0.0, round(gap, 1)))
        ppr = _ppr(paragraph)
        indent = item["x"] - left + float(item.get("left_indent") or 0)
        if indent:
            ppr.set("marL", str(int(Pt(indent))))
        if item.get("first_indent"):
            ppr.set("indent", str(int(Pt(float(item["first_indent"])))))
        if _is_bullet(item):
            # User 2026-10-12: "기사 요약 4불렛은 -를 일일이 넣지 말고 글머리 기호 활용해줘".
            etree.SubElement(ppr, qn("a:buFont")).set("typeface", FONT)
            etree.SubElement(ppr, qn("a:buChar")).set("char", "-")
        coloured = item.get("runs") or [(text, item["color"], bold) for text, bold in _pieces(item)]
        for text, color, bold in coloured:
            run = paragraph.add_run()
            run.text = text
            _font(run, size, color, bold or _bold_face(item["face"], float(item["size"])),
                  underline=bool(item.get("underline")))
            if item.get("url"):
                run.hyperlink.address = item["url"]
        item["_paragraph"] = paragraph
        bottom = item["baseline"] + (item["lines"] - 1) * lead + (lead - offset)
    return bottom


def _text_box(slide: Any, items: list[dict[str, Any]], page_width: float) -> Any:
    left = min(i["x"] for i in items)
    right = max(i["x"] + i["width"] for i in items)
    wrap = any(i["kind"] == "para" for i in items)
    width = min(right - left + SLACK, page_width - left - 2)
    top = items[0]["baseline"] - _metrics(items[0])[2]
    _, last_lead, last_offset = _metrics(items[-1])
    height = items[-1]["baseline"] + (items[-1]["lines"] - 1) * last_lead + (last_lead - last_offset) + 1 - top
    shape = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(max(width, 4)), Pt(max(height, 4)))
    _frame(shape.text_frame, wrap)
    _fill_paragraphs(shape.text_frame, items, left, top)
    for item in items:
        item["_shape"] = shape
    return shape


def _mergeable(prev: dict[str, Any], item: dict[str, Any]) -> bool:
    """A smaller line right under a text item (headline + meta) shares its text box."""
    if abs(prev["x"] - item["x"]) > 0.6 or item["size"] >= prev["size"]:
        return False
    last = prev["baseline"] + (prev["lines"] - 1) * prev["leading"]
    return bool(0 < item["baseline"] - last <= 1.8 * prev["size"])


def _solid(shape: Any, color: str | None, line: str | None = None, width: float = 0.75) -> None:
    if color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False


def _rect(slide: Any, item: dict[str, Any]) -> Any:
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if item["kind"] == "roundrect" else (
        MSO_SHAPE.OVAL if item["kind"] == "circle" else MSO_SHAPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Pt(item["x"]), Pt(item["top"]), Pt(item["w"]), Pt(item["h"]))
    if item["kind"] == "roundrect":
        shape.adjustments[0] = min(0.5, item["radius"] / max(1.0, min(item["w"], item["h"])))
    _solid(shape, item["fill_color"] if item["fill"] else None,
           item["stroke_color"] if item["stroke"] and not item["fill"] else None, item["line_width"])
    return shape


def _line(slide: Any, item: dict[str, Any]) -> Any:
    weight = max(0.25, float(item["line_width"]))
    if abs(item["y1"] - item["y2"]) < 0.01 or abs(item["x1"] - item["x2"]) < 0.01:
        # Rules as thin rectangles render at their true weight and are easy to select.
        x0, x1 = sorted((item["x1"], item["x2"]))
        y0, y1 = sorted((item["y1"], item["y2"]))
        horizontal = abs(item["y1"] - item["y2"]) < 0.01
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Pt(x0 if horizontal else x0 - weight / 2), Pt(y0 - weight / 2 if horizontal else y0),
            Pt(x1 - x0 if horizontal else weight), Pt(weight if horizontal else y1 - y0))
        _solid(shape, item["stroke_color"])
        return shape
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(item["x1"]), Pt(item["y1"]),
                                           Pt(item["x2"]), Pt(item["y2"]))
    connector.line.color.rgb = RGBColor.from_string(item["stroke_color"])
    connector.line.width = Pt(weight)
    return connector


def _table(slide: Any, group: dict[str, Any], margin: float) -> Any:
    args = group["args"]
    headers, rows, widths = args["headers"], args["rows"], list(args["widths"])
    top, row_h = float(args["top"]), float(args["row_h"])
    dy = float(args.get("label_dy") or 7)  # label top below the row top (trend tables centre smaller rows)
    grid = [headers, *rows]
    shape = slide.shapes.add_table(len(grid), len(widths), Pt(margin), Pt(top), Pt(sum(widths)),
                                   Pt(row_h * len(grid)))
    table = shape.table
    table._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_TABLE
    table.first_row = table.horz_banding = False
    for j, width in enumerate(widths):
        table.columns[j].width = Pt(width)
    fills = [i for i in group["items"] if i["kind"] == "rect" and i["fill"]]
    labels = [i for i in group["items"] if i["kind"] == "label"]
    edges = [margin + sum(widths[:j]) for j in range(len(widths) + 1)]
    for r in range(len(grid)):
        table.rows[r].height = Pt(row_h)
        row_top = top + r * row_h
        fill = next((f["fill_color"] for f in fills if abs(f["top"] - row_top) < 0.5), "FFFFFF")
        for c in range(len(widths)):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            label = next((x for x in labels if abs(x["top"] - (row_top + dy)) < 0.6
                          and edges[c] <= x["x"] + x["width"] / 2 < edges[c + 1]), None)
            if label:
                size, lead, offset = _metrics(label)
                pdf_size = float(label["size"])
            else:
                pdf_size = size = float(args.get("font", 8.8))
                lead, offset = _lead(1.3 * size), 0.755 * _lead(1.3 * size)
            cell.margin_left, cell.margin_right = Pt(5), Pt(2)
            cell.margin_top = Pt(max(0.0, dy + pdf_size - offset))
            cell.margin_bottom = Pt(0)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.line_spacing = Pt(lead)
            _ppr(paragraph)
            # A label centred in its PDF cell stays centred in the PowerPoint cell.
            if label and abs(label["x"] + label["width"] / 2 - (edges[c] + edges[c + 1]) / 2) < 1.5:
                paragraph.alignment = PP_ALIGN.CENTER
                cell.margin_left = cell.margin_right = Pt(2)
            run = paragraph.add_run()
            run.text = label["text"] if label else str(grid[r][c])
            _font(run, size, label["color"] if label else "26364B",
                  _bold_face(label["face"], pdf_size) if label else False)
    return shape


def _join_lines(labels: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Labels drawn side by side on one baseline (a coloured change inside a note) form one line."""
    lines: list[list[dict[str, Any]]] = []
    for label in labels:
        if lines and abs(lines[-1][0]["baseline"] - label["baseline"]) < 0.2:
            lines[-1].append(label)
        else:
            lines.append([label])
    joined = []
    for parts in lines:
        if len(parts) == 1:
            joined.append(parts[0])
            continue
        parts.sort(key=lambda x: x["x"])
        first = dict(parts[0])
        first.update(text="".join(x["text"] for x in parts),
                     width=parts[-1]["x"] + parts[-1]["width"] - parts[0]["x"],
                     runs=[(x["text"], x["color"], False) for x in parts])
        joined.append(first)
    return joined


def _stat_boxes(slide: Any, group: dict[str, Any]) -> None:
    boxes = [i for i in group["items"] if i["kind"] == "rect" and i["fill"]]
    labels = [i for i in group["items"] if i["kind"] == "label"]
    for box in boxes:
        inside = sorted((x for x in labels if box["x"] <= x["x"] <= box["x"] + box["w"]
                         and box["top"] <= x["top"] <= box["top"] + box["h"]), key=lambda x: x["baseline"])
        shape = _rect(slide, box)
        frame = shape.text_frame
        _frame(frame, False)
        if not inside:
            continue
        inside = _join_lines(inside)
        frame.margin_left = Pt(inside[0]["x"] - box["x"])
        first_top = inside[0]["baseline"] - _metrics(inside[0])[2]
        frame.margin_top = Pt(max(0.0, first_top - box["top"]))
        _fill_paragraphs(frame, inside, inside[0]["x"], first_top)


def _card(slide: Any, group: dict[str, Any], links: list[dict[str, Any]]) -> None:
    """Card background shape carries title, bullets and the source link in one text frame."""
    rects = [i for i in group["items"] if i["kind"] == "rect" and i["fill"]]
    texts = sorted((i for i in group["items"] if i["kind"] in ("label", "para")), key=lambda i: i["baseline"])
    lines = [i for i in group["items"] if i["kind"] == "line"]
    for link in (i for i in group["items"] if i["kind"] == "link"):
        for text in texts:
            if _overlap(link["box"], _bbox(text)) > 0.3:
                text["url"] = link["url"]
                text["underline"] = any(abs(ln["y1"] - (text["baseline"] + 2)) < 2.5 for ln in lines)
                link["_used"] = True
    links.extend(i for i in group["items"] if i["kind"] == "link" and not i.get("_used"))
    card = max(rects, key=lambda r: r["w"] * r["h"])
    shape = _rect(slide, card)
    # User 2026-10-30: PowerPoint "PDF로 저장" keeps shape click links but drops text-run links,
    # so the card shape also carries its source link (the run link stays for clicks in the deck).
    url = next((t["url"] for t in texts if t.get("url")), None)
    if url:
        shape.click_action.hyperlink.address = url
    for other in rects:
        if other is not card:
            _rect(slide, other)
    if not texts:
        return
    frame = shape.text_frame
    _frame(frame, True)
    left = min(t["x"] for t in texts)
    first_top = texts[0]["baseline"] - _metrics(texts[0])[2]
    frame.margin_left = Pt(left - card["x"])
    frame.margin_right = Pt(max(0.0, card["x"] + card["w"] - max(t["x"] + t["width"] for t in texts) - SLACK))
    frame.margin_top = Pt(max(0.0, first_top - card["top"]))
    _fill_paragraphs(frame, texts, left, first_top)
    for text in texts:
        text["_shape"] = shape


def _bbox(item: dict[str, Any]) -> tuple[float, float, float, float]:
    lead = item["leading"]
    top = item["baseline"] - item["size"]
    bottom = item["baseline"] + (item["lines"] - 1) * lead + 0.25 * item["size"]
    return item["x"], top, item["x"] + item["width"], bottom


def _overlap(box: Sequence[float], text: Sequence[float]) -> float:
    w = max(0.0, min(box[2], text[2]) - max(box[0], text[0]))
    h = max(0.0, min(box[3], text[3]) - max(box[1], text[1]))
    area = max(1e-6, (text[2] - text[0]) * (text[3] - text[1]))
    return w * h / area


def _chart(slide: Any, group: dict[str, Any], page_image: Any, scale: float, out: Path) -> None:
    a = group["args"]
    x0, y0 = float(a["x"]) - 4, float(a["top"]) - 1.5
    x1, y1 = float(a["x"]) + float(a["width"]) + 4, float(a["top"]) + float(a["height"]) + 2
    crop = page_image.crop((int(x0 * scale), int(y0 * scale), int(x1 * scale), int(y1 * scale))).convert("RGBA")
    # White page background becomes transparent so the picture never hides a neighbouring heading.
    crop.putdata([(r, g, b, 0) if min(r, g, b) >= 250 else (r, g, b, a) for r, g, b, a in crop.getdata()])
    out.parent.mkdir(parents=True, exist_ok=True)
    crop.save(out)
    slide.shapes.add_picture(str(out), Pt(x0), Pt(y0), Pt(x1 - x0), Pt(y1 - y0))


def _flush(slide: Any, pending: list[dict[str, Any]], texts: list[dict[str, Any]],
           counts: dict[str, int], page_width: float) -> None:
    """Write the waiting text items as one text box."""
    if pending:
        _text_box(slide, list(pending), page_width)
        texts.extend(pending)
        counts["text_boxes"] += 1
        pending.clear()


def write_deck(rec: Recorder, pdf_path: Path, dest: Path, *, page_width: float, margin: float,
               chart_dir: Path, scale: float = 3.0) -> dict[str, Any]:
    import pypdfium2 as pdfium  # type: ignore[import-untyped]

    if dest.exists():
        raise FileExistsError(dest)
    document = pdfium.PdfDocument(str(pdf_path))
    if len(document) != len(rec.pages):
        raise ValueError(f"recorded {len(rec.pages)} pages, PDF has {len(document)}")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(page_width), Pt(rec.H)
    slides = [deck.slides.add_slide(deck.slide_layouts[6]) for _ in rec.pages]
    counts = {"text_boxes": 0, "tables": 0, "cards": 0, "stat_boxes": 0, "charts": 0, "links": 0}
    for number, (slide, items) in enumerate(zip(slides, rec.pages, strict=True), start=1):
        page_image = None
        links: list[dict[str, Any]] = []
        jumps: list[dict[str, Any]] = []
        texts: list[dict[str, Any]] = []
        pending: list[dict[str, Any]] = []
        for item in items:
            kind = item["kind"]
            if kind in ("label", "para"):
                if pending and not _mergeable(pending[-1], item):
                    _flush(slide, pending, texts, counts, page_width)
                pending.append(item)
                continue
            _flush(slide, pending, texts, counts, page_width)
            if kind in ("rect", "roundrect", "circle"):
                if item["w"] > 0 and item["h"] > 0:
                    _rect(slide, item)
            elif kind == "line":
                _line(slide, item)
            elif kind == "image":
                slide.shapes.add_picture(item["path"], Pt(item["x"]), Pt(item["top"]), Pt(item["w"]), Pt(item["h"]))
            elif kind == "link":
                links.append(item)
            elif kind == "jump":
                jumps.append(item)
            elif kind == "table":
                _table(slide, item, margin)
                counts["tables"] += 1
            elif kind == "stats":
                _stat_boxes(slide, item)
                counts["stat_boxes"] += 1
            elif kind == "card":
                _card(slide, item, links)
                texts.extend(i for i in item["items"] if i["kind"] in ("label", "para"))
                counts["cards"] += 1
            elif kind == "chart":
                if page_image is None:
                    page_image = document[number - 1].render(scale=scale).to_pil()
                _chart(slide, item, page_image, scale, chart_dir / f"{dest.stem}_p{number:02d}_chart{counts['charts']}.png")
                counts["charts"] += 1
        _flush(slide, pending, texts, counts, page_width)
        for link in links:
            # Shape-level links keep the PDF colours (PowerPoint restyles hyperlinked runs).
            shapes = {id(t["_shape"]): t["_shape"] for t in texts
                      if "_shape" in t and _overlap(link["box"], _bbox(t)) > 0.3}
            free = [s for s in shapes.values() if not s.click_action.hyperlink.address]
            if free:
                for shape in free:
                    shape.click_action.hyperlink.address = link["url"]
            else:  # second link over the same text, or no recorded text: transparent clickable area
                x0, y0, x1, y1 = link["box"]
                area = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x0), Pt(y0), Pt(x1 - x0), Pt(y1 - y0))
                _solid(area, None)
                area.click_action.hyperlink.address = link["url"]
            counts["links"] += 1
        for jump in jumps:
            target = int(str(jump["dest"]).lstrip("p")) - 1
            for text in texts:
                if _overlap(jump["box"], _bbox(text)) > 0.3 and 0 <= target < len(slides):
                    text["_shape"].click_action.target_slide = slides[target]
    deck.save(str(dest))
    counts["slides"] = len(slides)
    return counts
